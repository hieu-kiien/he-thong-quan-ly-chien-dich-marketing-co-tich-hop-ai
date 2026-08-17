"""Document retrieval for the BAI03 project.

The module deliberately stops at retrieval and cited context. A provider-specific
LLM answer layer can consume ``KnowledgeService.search`` later without changing
the corpus or metadata contract.
"""

from __future__ import annotations

import hashlib
import re
from zipfile import BadZipFile
from pathlib import Path

from django.conf import settings


PROJECT_ROOT = Path(__file__).resolve().parents[3]
WORKSPACE_ROOT = PROJECT_ROOT.parent
Bai02_DOCUMENT_ROOT = WORKSPACE_ROOT / "Bai 02" / "CacGiaiDoanThucHien"
SUPPORTED_EXTENSIONS = {".md", ".yaml", ".yml", ".py", ".docx", ".pdf"}
SKIP_PARTS = {
    ".git",
    ".gitnexus",
    ".venv",
    ".venv-rag",
    "venv",
    "__pycache__",
    ".pytest_cache",
    ".ipynb_checkpoints",
    "bk",
    "mau",
    "slide_pdf",
    "qlbh demo",
}
FRONTMATTER_RE = re.compile(r"\A---\s*\n(.*?)\n---\s*(?:\n|\Z)", re.DOTALL)
TOKEN_RE = re.compile(r"[\wÀ-ỹ]+", re.UNICODE)
HEADING_RE = re.compile(r"^#{1,6}\s+(.+?)\s*$", re.MULTILINE)


class KnowledgeDependencyError(RuntimeError):
    """Raised when the optional vector/embedding runtime is not available."""


class KnowledgeSourceError(RuntimeError):
    """Raised when one source file cannot be parsed safely."""


def _normalise_path(path: Path | str) -> Path:
    return Path(str(path).replace("\\", "/"))


def source_path(path: Path | str) -> str:
    """Return a portable path suitable for citations and metadata."""

    candidate = _normalise_path(path)
    if candidate.is_absolute():
        try:
            candidate = candidate.resolve().relative_to(WORKSPACE_ROOT.resolve())
        except ValueError:
            pass
    return candidate.as_posix().lstrip("./")


def source_type(path: Path | str) -> str:
    portable = source_path(path).lower()
    suffix = _normalise_path(path).suffix.lower()
    if "/prompts/" in f"/{portable}/":
        return "prompt"
    if suffix == ".py":
        return "code"
    if suffix in {".yaml", ".yml"}:
        return "config"
    if suffix in {".docx", ".pdf"}:
        return "document"
    return "documentation"


def infer_status(path: Path | str) -> str:
    portable = source_path(path).lower()
    if "legacy" in portable:
        return "LEGACY"
    if "template" in portable or "/mau/" in f"/{portable}/":
        return "TEMPLATE"
    if "contextproject" in portable:
        return "DERIVED"
    if "/sales_management/" in f"/{portable}/" and portable.endswith(".py"):
        return "IMPLEMENTED"
    if "/prompts/" in f"/{portable}/":
        return "REFERENCE"
    return "REFERENCE"


def _frontmatter(text: str) -> dict[str, str]:
    match = FRONTMATTER_RE.match(text.replace("\r\n", "\n"))
    if not match:
        return {}
    values: dict[str, str] = {}
    for line in match.group(1).splitlines():
        key, separator, value = line.partition(":")
        if separator:
            values[key.strip()] = value.strip().strip("'\"")
    return values


def extract_metadata(path: Path | str, text: str) -> dict[str, str]:
    portable = source_path(path)
    parts = portable.lower().split("/")
    project_id = "BAI02-SDLC" if "bai 02" in parts else "BAI03-SALES-AI"
    values = _frontmatter(text)
    return {
        "source_path": portable,
        "source_type": source_type(path),
        "document_id": values.get("document_id") or portable,
        "project_id": values.get("project_id") or project_id,
        "status": (values.get("status") or infer_status(path)).upper(),
        "title": values.get("title") or _normalise_path(path).stem,
    }


def is_allowed_source(path: Path | str, include_code: bool = False) -> bool:
    candidate = _normalise_path(path)
    portable = source_path(candidate)
    lower = portable.lower()
    parts = {part.casefold() for part in candidate.parts}
    suffix = candidate.suffix.lower()

    if parts & SKIP_PARTS or "legacy-prime-number-utilities" in lower:
        return False
    if candidate.name.startswith("~$"):
        return False
    if suffix not in SUPPORTED_EXTENSIONS:
        return False
    if "bai 02" in parts and "cacgiaidoanthuchien" not in lower:
        return False
    if suffix == ".py":
        if not include_code:
            return False
        if not ({"sales_management", "code qlbh"} & parts):
            return False
    if suffix in {".docx", ".pdf"} and "docs" not in parts and "cacgiaidoanthuchien" not in lower:
        return False
    if "code qlbh" in parts and not include_code:
        return False
    return True


def iter_source_files(include_code: bool = False) -> list[Path]:
    files: list[Path] = []
    roots = [PROJECT_ROOT, Bai02_DOCUMENT_ROOT]
    for root in roots:
        if not root.exists():
            continue
        for candidate in root.rglob("*"):
            if candidate.is_file() and is_allowed_source(candidate, include_code=include_code):
                files.append(candidate)
    return sorted(set(files), key=lambda item: source_path(item).lower())


def chunk_text(text: str, max_chars: int = 1800, overlap: int = 250) -> list[str]:
    """Split text into bounded, overlapping chunks without empty records."""

    if max_chars <= 0 or overlap < 0 or overlap >= max_chars:
        raise ValueError("overlap must be smaller than max_chars")
    clean = re.sub(r"\r\n?", "\n", text).strip()
    if not clean:
        return []

    paragraphs = [item.strip() for item in re.split(r"\n\s*\n", clean) if item.strip()]
    chunks: list[str] = []
    current = ""

    def flush(value: str) -> str:
        if value:
            chunks.append(value)
        return value[-overlap:] if overlap else ""

    for paragraph in paragraphs:
        while len(paragraph) > max_chars:
            piece = paragraph[:max_chars]
            if current:
                current = flush(current)
            current = flush(piece)
            paragraph = paragraph[max_chars - overlap :]
        candidate = f"{current}\n\n{paragraph}" if current else paragraph
        if len(candidate) <= max_chars:
            current = candidate
        else:
            current = flush(current)
            current = f"{current}\n\n{paragraph}" if current else paragraph
    if current:
        chunks.append(current)
    return chunks


def stable_chunk_id(path: str, content_hash: str, chunk_index: int) -> str:
    value = f"{source_path(path)}:{content_hash}:{chunk_index}"
    return hashlib.sha256(value.encode("utf-8")).hexdigest()[:32]


def _read_source(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        try:
            from docx import Document
        except ImportError as exc:
            raise KnowledgeDependencyError("Cài python-docx để đọc DOCX") from exc
        try:
            document = Document(path)
        except (BadZipFile, OSError, ValueError) as exc:
            raise KnowledgeSourceError(f"DOCX không hợp lệ: {source_path(path)}") from exc
        blocks = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    blocks.append("TABLE | " + " | ".join(cells))
        return "\n\n".join(blocks)
    if suffix == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise KnowledgeDependencyError("Cài pypdf để đọc PDF") from exc
        try:
            reader = PdfReader(path)
            pages = []
            for index, page in enumerate(reader.pages, start=1):
                pages.append(f"[Page {index}]\n{page.extract_text() or ''}")
        except (OSError, ValueError) as exc:
            raise KnowledgeSourceError(f"PDF không đọc được: {source_path(path)}") from exc
        return "\n\n".join(pages)
    return path.read_text(encoding="utf-8", errors="replace")


def _section(text: str) -> str:
    match = HEADING_RE.search(text)
    return match.group(1).strip() if match else ""


def format_context(results: list[dict]) -> str:
    blocks = []
    for index, result in enumerate(results, start=1):
        metadata = result.get("metadata", {})
        citation = metadata.get("source_path", "unknown")
        if metadata.get("section"):
            citation += f"#{metadata['section']}"
        blocks.append(f"[{index}] {citation}\n{result.get('text', '').strip()}")
    return "\n\n".join(blocks)


class KnowledgeService:
    """Persistent Chroma retrieval with multilingual sentence embeddings."""

    def __init__(self, store_path: str | Path | None = None, model_name: str | None = None):
        self.store_path = Path(
            store_path
            or getattr(settings, "KNOWLEDGE_STORE_PATH", PROJECT_ROOT / "knowledge_store")
        )
        self.model_name = model_name or getattr(
            settings, "RAG_EMBEDDING_MODEL", "intfloat/multilingual-e5-small"
        )
        self.collection_name = getattr(settings, "RAG_COLLECTION_NAME", "bai03_knowledge")
        self._client = None
        self._collection = None
        self._model = None

    def _ensure_runtime(self):
        if self._collection is not None:
            return
        try:
            import chromadb
            from sentence_transformers import SentenceTransformer
        except ImportError as exc:
            raise KnowledgeDependencyError(
                "Cài chromadb và sentence-transformers trong môi trường dự án"
            ) from exc
        self.store_path.mkdir(parents=True, exist_ok=True)
        self._client = chromadb.PersistentClient(path=str(self.store_path))
        self._collection = self._client.get_or_create_collection(
            name=self.collection_name,
            metadata={"hnsw:space": "cosine"},
        )
        self._model = SentenceTransformer(self.model_name)

    def _encode(self, texts: list[str], query: bool = False) -> list[list[float]]:
        self._ensure_runtime()
        prefix = "query: " if query else "passage: "
        vectors = self._model.encode(
            [prefix + text for text in texts],
            normalize_embeddings=True,
            show_progress_bar=False,
        )
        return [vector.tolist() if hasattr(vector, "tolist") else list(vector) for vector in vectors]

    def index(
        self,
        *,
        include_code: bool = False,
        rebuild: bool = False,
        max_chars: int = 1800,
        overlap: int = 250,
    ) -> dict[str, int | str]:
        self._ensure_runtime()
        if rebuild:
            self._client.delete_collection(name=self.collection_name)
            self._collection = None
            self._ensure_runtime()

        files_indexed = 0
        chunks_indexed = 0
        files_skipped = 0
        for path in iter_source_files(include_code=include_code):
            try:
                text = _read_source(path)
            except KnowledgeDependencyError:
                raise
            except KnowledgeSourceError:
                files_skipped += 1
                continue
            chunks = chunk_text(text, max_chars=max_chars, overlap=overlap)
            if not chunks:
                continue
            content_hash = hashlib.sha256(path.read_bytes()).hexdigest()
            base_metadata = extract_metadata(path, text)
            base_metadata.update(
                {
                    "content_hash": content_hash,
                    "last_modified": path.stat().st_mtime_ns,
                }
            )
            ids = [stable_chunk_id(base_metadata["source_path"], content_hash, i) for i in range(len(chunks))]
            metadatas = [
                {
                    **base_metadata,
                    "chunk_index": i,
                    "section": _section(chunk),
                    "char_count": len(chunk),
                }
                for i, chunk in enumerate(chunks)
            ]
            for start in range(0, len(chunks), 32):
                end = start + 32
                self._collection.upsert(
                    ids=ids[start:end],
                    documents=chunks[start:end],
                    metadatas=metadatas[start:end],
                    embeddings=self._encode(chunks[start:end]),
                )
            files_indexed += 1
            chunks_indexed += len(chunks)
        return {
            "files": files_indexed,
            "chunks": chunks_indexed,
            "skipped": files_skipped,
            "collection_count": self._collection.count(),
            "store_path": str(self.store_path),
        }

    def search(
        self,
        query: str,
        *,
        top_k: int = 5,
        status: str | None = None,
        source_type_filter: str | None = None,
    ) -> list[dict]:
        if not query or not query.strip():
            raise ValueError("query không được để trống")
        if top_k < 1 or top_k > 50:
            raise ValueError("top_k phải nằm trong khoảng 1..50")
        self._ensure_runtime()
        count = self._collection.count()
        if not count:
            return []
        filters = []
        if status:
            filters.append({"status": status.upper()})
        if source_type_filter:
            filters.append({"source_type": source_type_filter})
        where = None
        if len(filters) == 1:
            where = filters[0]
        elif filters:
            where = {"$and": filters}
        result = self._collection.query(
            query_embeddings=self._encode([query], query=True),
            n_results=min(count, max(top_k * 4, top_k)),
            where=where,
            include=["documents", "metadatas", "distances"],
        )
        query_tokens = set(TOKEN_RE.findall(query.casefold()))
        rows = []
        for index, text in enumerate(result.get("documents", [[]])[0]):
            metadata = result.get("metadatas", [[]])[0][index]
            distance = float(result.get("distances", [[]])[0][index])
            tokens = set(TOKEN_RE.findall(text.casefold()))
            lexical = len(query_tokens & tokens) / max(len(query_tokens), 1)
            dense = max(0.0, min(1.0, 1.0 - distance))
            rows.append(
                {
                    "id": result.get("ids", [[]])[0][index],
                    "text": text,
                    "metadata": metadata,
                    "distance": distance,
                    "score": round(0.75 * dense + 0.25 * lexical, 6),
                }
            )
        rows.sort(key=lambda item: item["score"], reverse=True)
        return rows[:top_k]
