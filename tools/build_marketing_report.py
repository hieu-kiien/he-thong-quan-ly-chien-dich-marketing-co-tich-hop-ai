"""Build the AIA331 marketing report from the canonical Markdown source."""

from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "submission" / "Nhom25" / "Bao_cao_du_an_marketing_ai.md"
OUTPUT = ROOT / "submission" / "Nhom25" / "Bao_cao_du_an_marketing_ai.docx"
FONT = "Times New Roman"
BLUE = RGBColor(31, 78, 121)
MUTED = RGBColor(89, 89, 89)
TABLE_HEADER = "D9E2F3"


def set_run_font(run, size=13, bold=None, italic=None, color=None, name=FONT):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def set_cell_shading(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="B7C9D6", size="6"):
    properties = table._tbl.tblPr
    borders = properties.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_table_widths(table, widths_mm):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row in table.rows:
        for index, width in enumerate(widths_mm):
            row.cells[index].width = Mm(width)
            set_cell_margins(row.cells[index])
            row.cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, text, end])
    set_run_font(run, size=9, color=MUTED)


def add_toc_field(document):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = ' TOC \\o "1-2" \\h \\z \\u '
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "Mục lục sẽ được cập nhật khi mở trong Word."
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, placeholder, end])
    set_run_font(run, size=11, color=MUTED)


def add_static_toc(document):
    entries = [
        "1. Nguồn và phạm vi",
        "2. Phân tích bài toán quản lý",
        "3. Yêu cầu chức năng",
        "4. Yêu cầu phi chức năng",
        "5. Use case",
        "6. Thiết kế dữ liệu",
        "7. Kiến trúc hệ thống",
        "8. Định vị AI",
        "9. Prompt và AI contract",
        "10. Kiểm thử và minh chứng",
        "11. Rủi ro và điểm mở",
        "12. Kế hoạch triển khai tiếp",
        "Kết luận",
    ]
    for entry in entries:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Mm(8)
        paragraph.paragraph_format.space_after = Pt(5)
        add_rich_text(paragraph, entry, size=12)


def new_numbering_id(document):
    numbering = document.part.numbering_part.element
    abstract_ids = []
    for abstract in numbering.findall(qn("w:abstractNum")):
        for num_format in abstract.findall(".//" + qn("w:numFmt")):
            if num_format.get(qn("w:val")) == "decimal":
                abstract_ids.append(abstract.get(qn("w:abstractNumId")))
                break
    abstract_id = abstract_ids[0] if abstract_ids else "0"
    existing = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    num_id = max(existing or [0]) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), abstract_id)
    num.append(abstract_ref)
    level_override = OxmlElement("w:lvlOverride")
    level_override.set(qn("w:ilvl"), "0")
    start_override = OxmlElement("w:startOverride")
    start_override.set(qn("w:val"), "1")
    level_override.append(start_override)
    num.append(level_override)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id):
    properties = paragraph._p.get_or_add_pPr()
    num_pr = properties.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        properties.append(num_pr)
    ilvl = num_pr.find(qn("w:ilvl"))
    if ilvl is None:
        ilvl = OxmlElement("w:ilvl")
        num_pr.append(ilvl)
    ilvl.set(qn("w:val"), "0")
    num = num_pr.find(qn("w:numId"))
    if num is None:
        num = OxmlElement("w:numId")
        num_pr.append(num)
    num.set(qn("w:val"), str(num_id))


def repeat_table_header(row):
    properties = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    properties.append(header)


def add_rich_text(paragraph, text, size=13):
    pattern = re.compile(r"(\*\*.*?\*\*|`.*?`|\[.*?\]\(.*?\))")
    cursor = 0
    for match in pattern.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor : match.start()])
            set_run_font(run, size=size)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=size, bold=True)
        elif token.startswith("["):
            label = token[1 : token.find("]")]
            run = paragraph.add_run(label)
            set_run_font(run, size=size, color=BLUE)
        else:
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=max(9.5, size - 1), name="Consolas")
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_run_font(run, size=size)


def add_paragraph(document, text, style=None, size=13, italic=False, color=None):
    paragraph = document.add_paragraph(style=style)
    paragraph.paragraph_format.line_spacing = 1.3
    paragraph.paragraph_format.space_after = Pt(6)
    add_rich_text(paragraph, text, size=size)
    if italic:
        for run in paragraph.runs:
            run.italic = True
    if color:
        for run in paragraph.runs:
            run.font.color.rgb = color
    return paragraph


def add_markdown_table(document, lines):
    rows = []
    for line in lines:
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        if all(set(cell) <= {"-", ":", " "} for cell in cells):
            continue
        rows.append(cells)
    if not rows:
        return
    columns = max(len(row) for row in rows)
    rows = [row + [""] * (columns - len(row)) for row in rows]
    table = document.add_table(rows=len(rows), cols=columns)
    table.style = "Table Grid"
    set_table_borders(table)
    if columns == 2:
        widths = [45, 110]
    elif columns == 3:
        widths = [30, 60, 65]
    elif columns == 4:
        widths = [24, 44, 48, 39]
    elif columns == 5:
        widths = [20, 38, 35, 34, 28]
    else:
        widths = [155 / columns] * columns
    set_table_widths(table, widths)
    repeat_table_header(table.rows[0])
    for row_index, row in enumerate(rows):
        for col_index, value in enumerate(row):
            cell = table.cell(row_index, col_index)
            cell.text = ""
            if row_index == 0:
                set_cell_shading(cell, TABLE_HEADER)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.15
            add_rich_text(paragraph, value, size=9.5)
            if row_index == 0:
                for run in paragraph.runs:
                    run.bold = True
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def configure_document(document):
    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.left_margin = Mm(35)
    section.right_margin = Mm(20)
    section.top_margin = Mm(25)
    section.bottom_margin = Mm(25)
    section.header_distance = Mm(12)
    section.footer_distance = Mm(12)

    normal = document.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(13)
    normal.paragraph_format.line_spacing = 1.3
    normal.paragraph_format.space_after = Pt(6)

    for name, size, before, after in (("Heading 1", 16, 16, 8), ("Heading 2", 14, 12, 6), ("Heading 3", 13, 8, 4)):
        style = document.styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLUE
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.line_spacing = 1.15

    footer = section.footer
    footer_paragraph = footer.paragraphs[0]
    footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_paragraph.paragraph_format.space_before = Pt(0)
    footer_paragraph.paragraph_format.space_after = Pt(0)
    run = footer_paragraph.add_run("AIA331-80300-MARKETING-AI · Nhóm 25 · Trang ")
    set_run_font(run, size=9, color=MUTED)
    add_page_field(footer_paragraph)


def add_cover(document):
    for text, size, bold, space_after in (
        ("ĐẠI HỌC THÁI NGUYÊN", 13, False, 4),
        ("TRƯỜNG ĐẠI HỌC CÔNG NGHỆ THÔNG TIN VÀ TRUYỀN THÔNG", 13, True, 2),
        ("KHOA CÔNG NGHỆ THÔNG TIN", 13, True, 36),
    ):
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(space_after)
        run = paragraph.add_run(text)
        set_run_font(run, size=size, bold=bold)

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run("BÁO CÁO DỰ ÁN")
    set_run_font(run, size=20, bold=True, color=BLUE)

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run("HỌC PHẦN: ỨNG DỤNG TRÍ TUỆ NHÂN TẠO - AIA331")
    set_run_font(run, size=13, bold=True)

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(46)
    run = paragraph.add_run("HỆ THỐNG QUẢN LÝ CHIẾN DỊCH MARKETING\nCÓ TÍCH HỢP AI")
    set_run_font(run, size=18, bold=True, color=BLUE)

    table = document.add_table(rows=7, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table, color="FFFFFF", size="0")
    set_table_widths(table, [45, 110])
    metadata = [
        ("Mã số", "80300"),
        ("Hình thức", "Dự án"),
        ("Nhóm", "25"),
        ("Sinh viên", "Nguyễn Hải Đăng; Vũ Hiếu Kiên"),
        ("Mã sinh viên", "dtc2451200051; dtc245200244"),
        ("Lớp", "CNTTK23C"),
        ("Cấp ưu tiên", "P1 / HIGH — hồ sơ bám hai nguồn P0 / CRITICAL"),
    ]
    for row, (label, value) in zip(table.rows, metadata):
        for cell in row.cells:
            set_cell_margins(cell, top=80, bottom=80, start=80, end=80)
        row.cells[0].text = ""
        row.cells[1].text = ""
        p0 = row.cells[0].paragraphs[0]
        p1 = row.cells[1].paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        p1.paragraph_format.space_after = Pt(0)
        add_rich_text(p0, label, size=11.5)
        add_rich_text(p1, value, size=11.5)
        for run in p0.runs:
            run.bold = True

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(48)
    run = paragraph.add_run("Thái Nguyên, 2026")
    set_run_font(run, size=13)
    document.add_page_break()

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(12)
    run = paragraph.add_run("MỤC LỤC")
    set_run_font(run, size=16, bold=True, color=BLUE)
    add_static_toc(document)
    document.add_page_break()


def parse_report(document):
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    in_frontmatter = False
    frontmatter_seen = 0
    filtered = []
    for line in lines:
        if line.strip() == "---" and frontmatter_seen < 2:
            frontmatter_seen += 1
            in_frontmatter = frontmatter_seen == 1
            continue
        if in_frontmatter:
            continue
        filtered.append(line.rstrip())

    start = next((index for index, line in enumerate(filtered) if line.strip() == "## Tóm tắt"), 0)
    lines = filtered[start:]
    index = 0
    in_code = False
    code_lines = []
    numbered_active = False
    current_num_id = None
    while index < len(lines):
        line = lines[index]
        if line.strip() == "---":
            numbered_active = False
            index += 1
            continue
        if line.startswith("```"):
            if not in_code:
                in_code = True
                code_lines = []
            else:
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.left_indent = Mm(5)
                paragraph.paragraph_format.right_indent = Mm(5)
                paragraph.paragraph_format.space_after = Pt(8)
                paragraph.paragraph_format.line_spacing = 1.1
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "F2F4F7")
                paragraph._p.get_or_add_pPr().append(shading)
                run = paragraph.add_run("\n".join(code_lines))
                set_run_font(run, size=9.5, name="Consolas")
                in_code = False
            index += 1
            continue
        if in_code:
            code_lines.append(line)
            index += 1
            continue

        if line.startswith("|"):
            numbered_active = False
            table_lines = []
            while index < len(lines) and lines[index].startswith("|"):
                table_lines.append(lines[index])
                index += 1
            add_markdown_table(document, table_lines)
            continue
        if not line.strip():
            numbered_active = False
            index += 1
            continue
        heading = re.match(r"^(#{2,4})\s+(.*)$", line)
        if heading:
            level = len(heading.group(1)) - 1
            paragraph = document.add_paragraph(heading.group(2), style=f"Heading {min(level, 3)}")
            if heading.group(2).startswith("3. Yêu cầu chức năng"):
                paragraph.paragraph_format.page_break_before = True
            numbered_active = False
            index += 1
            continue
        bullet = re.match(r"^[-*]\s+(.*)$", line)
        if bullet:
            numbered_active = False
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.paragraph_format.line_spacing = 1.3
            paragraph.paragraph_format.space_after = Pt(4)
            add_rich_text(paragraph, bullet.group(1), size=13)
            index += 1
            continue
        numbered = re.match(r"^\d+\.\s+(.*)$", line)
        if numbered:
            paragraph = document.add_paragraph(style="List Number")
            paragraph.paragraph_format.line_spacing = 1.3
            paragraph.paragraph_format.space_after = Pt(4)
            if not numbered_active:
                current_num_id = new_numbering_id(document)
            apply_numbering(paragraph, current_num_id)
            add_rich_text(paragraph, numbered.group(1), size=13)
            numbered_active = True
            index += 1
            continue
        numbered_active = False
        add_paragraph(document, line)
        index += 1


def main():
    document = Document()
    configure_document(document)
    add_cover(document)
    parse_report(document)
    document.core_properties.title = "Báo cáo dự án — Hệ thống quản lý chiến dịch marketing có tích hợp AI"
    document.core_properties.subject = "AIA331 / 80300 / Nhóm 25"
    document.core_properties.author = "Nhóm 25 — CNTTK23C"
    document.core_properties.keywords = "AIA331, marketing campaign, AI, Nhóm 25"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print("created marketing report docx")


if __name__ == "__main__":
    main()
