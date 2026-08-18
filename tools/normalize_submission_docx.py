from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


FONT = "Times New Roman"
BODY_SIZE = 13
TABLE_SIZE = 10.5


def set_rfonts(rpr, name: str) -> None:
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), name)
    rfonts.set(qn("w:cs"), name)


def set_run_font(run, name: str = FONT, size: float = BODY_SIZE, bold=None, italic=None) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if run._element.rPr is None:
        run._element.get_or_add_rPr()
    set_rfonts(run._element.rPr, name)


def set_style(style, size: float, bold=False, italic=False, before=0, after=6, line=1.3) -> None:
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = italic
    if style._element.rPr is None:
        style._element.get_or_add_rPr()
    set_rfonts(style._element.rPr, FONT)
    pf = style.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line


def ensure_style(doc: Document, name: str):
    try:
        return doc.styles[name]
    except KeyError:
        return doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)


def clear_paragraph(paragraph) -> None:
    p = paragraph._p
    for child in list(p):
        if child.tag != qn("w:pPr"):
            p.remove(child)


def remove_paragraph(paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)


def set_paragraph_text(paragraph, text: str, style: str | None = None) -> None:
    clear_paragraph(paragraph)
    if style:
        paragraph.style = style
    run = paragraph.add_run(text)
    set_run_font(run)


def add_page_field(paragraph) -> None:
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run = OxmlElement("w:r")
    run.append(begin)
    run.append(instruction)
    run.append(separate)
    run.append(text)
    run.append(end)
    paragraph._p.append(run)


def add_toc_field(paragraph) -> None:
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), 'TOC \\o "1-3" \\h \\z \\u')
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    set_rfonts(rpr, FONT)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), str(BODY_SIZE * 2))
    rpr.append(size)
    run.append(rpr)
    field.append(run)
    paragraph._p.append(field)


def clear_container(container) -> None:
    for child in list(container._element):
        container._element.remove(child)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cover_paragraph(paragraph, size: float, bold=True, after=6) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.15
    for run in paragraph.runs:
        set_run_font(run, size=size, bold=bold)


def normalize(input_path: Path, output_path: Path) -> None:
    doc = Document(str(input_path))
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(3.5)
    section.right_margin = Cm(2.0)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.25)

    normal = doc.styles["Normal"]
    set_style(normal, BODY_SIZE, before=6, after=6, line=1.3)
    for name, size, before, after in (
        ("Heading 1", 16, 14, 8),
        ("Heading 2", 14, 11, 6),
        ("Heading 3", 13, 8, 4),
    ):
        set_style(doc.styles[name], size, bold=True, before=before, after=after, line=1.15)
        doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
        doc.styles[name].paragraph_format.keep_with_next = True

    title = ensure_style(doc, "Title")
    set_style(title, 20, bold=True, after=8, line=1.15)
    title.font.color.rgb = RGBColor(0, 0, 0)
    subtitle = ensure_style(doc, "Subtitle")
    set_style(subtitle, 15, bold=True, after=12, line=1.15)
    subtitle.font.color.rgb = RGBColor(0, 0, 0)
    caption = ensure_style(doc, "Caption")
    set_style(caption, 11, italic=True, after=6, line=1.15)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.keep_with_next = True
    code_style = ensure_style(doc, "Code Block")
    code_style.font.name = "Consolas"
    code_style.font.size = Pt(8)
    code_style.paragraph_format.space_before = Pt(0)
    code_style.paragraph_format.space_after = Pt(0)
    code_style.paragraph_format.line_spacing = 1.0
    for list_name in ("List Bullet", "List Number"):
        try:
            set_style(doc.styles[list_name], BODY_SIZE, after=3, line=1.3)
        except KeyError:
            pass

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text.startswith("Hình ") or text.startswith("Bảng "):
            paragraph.style = "Caption"
        elif paragraph.style and paragraph.style.name == "Code Block":
            for run in paragraph.runs:
                set_run_font(run, name="Consolas", size=8)
        else:
            for run in paragraph.runs:
                set_run_font(run)
        if text.startswith("Phụ lục") or text.startswith("TÀI LIỆU THAM KHẢO"):
            paragraph.paragraph_format.page_break_before = True

    # The first three existing paragraphs are the report title block.
    first = doc.paragraphs[0]
    first.insert_paragraph_before("ĐẠI HỌC THÁI NGUYÊN")
    first.insert_paragraph_before("TRƯỜNG ĐẠI HỌC CÔNG NGHỆ THÔNG TIN VÀ TRUYỀN THÔNG")
    first.insert_paragraph_before("KHOA CÔNG NGHỆ THÔNG TIN")
    cover_paras = doc.paragraphs[:6]
    for paragraph in cover_paras[:3]:
        set_cover_paragraph(paragraph, 12, bold=True, after=2)
    set_cover_paragraph(cover_paras[3], 20, bold=True, after=8)
    set_cover_paragraph(cover_paras[4], 16, bold=True, after=8)
    set_cover_paragraph(cover_paras[5], 11, bold=False, after=12)
    cover_paras[5].paragraph_format.keep_with_next = True

    # Make the metadata table explicit and add the course/teacher fields.
    metadata = doc.tables[0]
    rows = [
        ("Bài kiểm tra", "Bài kiểm tra thường xuyên 1 — Bài kiểm tra số 01"),
        ("Học phần", "Ứng dụng AI"),
        ("Giảng viên", "Nguyễn Tuấn Anh"),
        ("Mã dự án", "BAI03-SALES-AI"),
        ("Cấp ưu tiên", "P1 / HIGH — ưu tiên cao nhất cho hồ sơ BAI03 và truy hồi AI"),
        ("Nhóm", "Nhóm 25"),
        ("Thành viên", "Nguyễn Hải Đăng (Thành viên); Vũ Hiếu Kiên (Thành viên)"),
        ("Mã sinh viên", "dtc2451200051; dtc245200244"),
        ("Lớp", "CNTTK23C"),
        ("Đơn vị", "Trường Đại học Công nghệ Thông tin và Truyền thông — Khoa Công nghệ thông tin"),
        ("Ngày nộp", "18/08/2026"),
        ("Trạng thái", "Bản nộp đã kiểm tra chéo; phần chưa triển khai ghi PROPOSED/OPEN"),
    ]
    for index, (label, value) in enumerate(rows):
        if index < len(metadata.rows):
            row = metadata.rows[index]
        else:
            row = metadata.add_row()
        row.cells[0].text = label
        row.cells[1].text = value
        for cell_index, cell in enumerate(row.cells):
            set_cell_margins(cell)
            if cell_index == 0:
                set_cell_shading(cell, "E8EEF5")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.line_spacing = 1.1
                for run in paragraph.runs:
                    set_run_font(run, size=10.5, bold=cell_index == 0)
    metadata.style = "Table Grid"

    # Replace the static TOC list with a real Word TOC field.
    toc_heading = next((p for p in doc.paragraphs if p.text.strip().startswith("Mục lục")), None)
    first_content = next(
        (
            p
            for p in doc.paragraphs
            if p.text.strip() == "0. Quy ước đọc báo cáo và quyết định kiểm soát"
            and p.style
            and p.style.name.startswith("Heading")
        ),
        None,
    )
    if toc_heading is not None and first_content is not None:
        set_paragraph_text(toc_heading, "MỤC LỤC", "Heading 1")
        toc_heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        toc_paragraphs = []
        seen_toc = False
        for paragraph in list(doc.paragraphs):
            if paragraph._p is toc_heading._p:
                seen_toc = True
                continue
            if seen_toc and paragraph._p is first_content._p:
                break
            if seen_toc:
                toc_paragraphs.append(paragraph)
        for paragraph in toc_paragraphs:
            remove_paragraph(paragraph)
        toc_field_paragraph = first_content.insert_paragraph_before()
        toc_field_paragraph.paragraph_format.space_after = Pt(12)
        add_toc_field(toc_field_paragraph)
        first_content.paragraph_format.page_break_before = True

    # Keep the report language final rather than draft-like.
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == "Bản thảo mục 1–3 — Phân tích và yêu cầu hệ thống":
            set_paragraph_text(paragraph, "Mở đầu và phạm vi báo cáo", "Heading 1")
        if "Phạm vi bản thảo:" in paragraph.text:
            paragraph.text = paragraph.text.replace("Phạm vi bản thảo:", "Phạm vi báo cáo:")
            for run in paragraph.runs:
                set_run_font(run)

    # Add the new reference section before Appendix A if the source DOCX predates it.
    appendix = next(
        (
            p
            for p in doc.paragraphs
            if p.text.strip().startswith("Phụ lục A")
            and p.style
            and p.style.name.startswith("Heading")
        ),
        None,
    )
    if appendix is not None and not any(p.text.strip() == "TÀI LIỆU THAM KHẢO" for p in doc.paragraphs):
        references = [
            "TÀI LIỆU THAM KHẢO",
            "1. Nhóm 25, project.md — yêu cầu gốc và rubric của BAI03-SALES-AI.",
            "2. Nhóm 25, informember.md — thông tin nhóm, lớp, trường và khoa.",
            "3. Nhóm 25, docs/01-requirements-summary.md — ma trận yêu cầu chuẩn hóa.",
            "4. Nhóm 25, docs/02-architecture-and-code-status.md — đối chiếu kiến trúc và hiện trạng code.",
            "5. Nhóm 25, docs/03-business-flows.md và docs/04-ai-specification.md — nghiệp vụ, đặc tả AI và guardrail.",
            "6. Khoa CNTT ICTU, mẫu quyển báo cáo: https://fit.ictu.edu.vn/nhiem-vu-cua-gvql-truong-doan-va-mau-quyen-bao-cao/",
            "7. Khoa CNTT ICTU, Quyết định 147: https://fit.ictu.edu.vn/quyet-dinh-so-147-ban-hanh-quy-trinh-thuc-hien-do-an-khoa-luan-tot-nghiep-nam-2025/",
        ]
        for index, text in enumerate(references):
            paragraph = appendix.insert_paragraph_before(text)
            paragraph.paragraph_format.space_after = Pt(6)
            if index == 0:
                paragraph.style = "Heading 1"
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                set_run_font(run, size=16 if index == 0 else BODY_SIZE, bold=index == 0)
        appendix.paragraph_format.page_break_before = True
        paragraphs_after_insert = list(doc.paragraphs)
        appendix_index = next(
            (index for index, paragraph in enumerate(paragraphs_after_insert) if paragraph._p is appendix._p),
            None,
        )
        if appendix_index is not None and appendix_index + 1 < len(paragraphs_after_insert):
            # The source DOCX contains a second break before the first appendix
            # paragraph; remove it so the appendix heading is not stranded alone.
            paragraphs_after_insert[appendix_index + 1].paragraph_format.page_break_before = False

    # Avoid a running header; retain only a centered page number in the footer.
    clear_container(section.header)
    header_paragraph = section.header.add_paragraph()
    header_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    clear_container(section.footer)
    footer_paragraph = section.footer.add_paragraph()
    footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_paragraph.paragraph_format.space_before = Pt(0)
    footer_paragraph.paragraph_format.space_after = Pt(0)
    run = footer_paragraph.add_run("Trang ")
    set_run_font(run, size=10)
    add_page_field(footer_paragraph)

    for shape in doc.inline_shapes:
        if shape.width > Cm(15.0):
            ratio = shape.height / shape.width
            shape.width = Cm(15.0)
            shape.height = int(shape.width * ratio)

    doc.core_properties.title = "Báo cáo phân tích và thiết kế — BAI03-SALES-AI"
    doc.core_properties.author = "Nhóm 25 — Nguyễn Hải Đăng; Vũ Hiếu Kiên"
    doc.core_properties.subject = "Báo cáo học phần Ứng dụng AI"
    doc.core_properties.comments = "Nguồn canonical: Bao_cao_phan_tich_thiet_ke_BAI03.md; version 1.1"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    normalize(args.input, args.output)
    print(f"wrote {args.output}")


if __name__ == "__main__":
    main()
